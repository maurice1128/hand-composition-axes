"""Render docs/tmlr/PAPER_TMLR.md as a formatted reading copy (.docx, and .pdf through Word): journal-style title,
black Times headings, three-line tables with captions above, figure captions below, numbered citations resolved
from references.bib, page numbers.

Usage: python scripts/tmlr_md_to_docx.py OUT.docx
"""
import re
import unicodedata
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "tmlr"
md = (SRC / "PAPER_TMLR.md").read_text(encoding="utf-8")
md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
bib = (SRC / "references.bib").read_text(encoding="utf-8")
FONT = "Times New Roman"
# --named: the public copy for the project page (author shown); default: the anonymous review copy.
# --applicant: the named copy for application documents (name and department as the author asked, 2026-10-02).
APPLICANT = "--applicant" in sys.argv
NAMED = "--named" in sys.argv or APPLICANT
AUTHOR = ("WANG MU HUA" if APPLICANT else
          "Mu-Hua (Maurice) Wang, National Yang Ming Chiao Tung University" if NAMED else "Anonymous authors")
AFFIL = ("Department of Biomedical Engineering, National Yang Ming Chiao Tung University (sole author)"
         if APPLICANT else None)
if NAMED:  # one named author: the AI-assistance declaration speaks of "the author", not "the authors"
    md = md.replace("The authors set the research questions, decided which claims\nto make and are responsible",
                    "The author set the research questions, decided which claims\nto make and is responsible")
STATUS = ("Manuscript prepared for Transactions on Machine Learning Research (not yet submitted)" if NAMED
          else "Paper under double-blind review")

entries = {}
for m in re.finditer(r"@(\w+)\{([^,]+),(.*?)\n\}", bib, re.S):
    fields = {}
    for f in re.finditer(r"(\w+)\s*=\s*\{((?:[^{}]|\{[^{}]*\})*)\}", m.group(3), re.S):
        val = " ".join(f.group(2).split())
        # LaTeX accents ({\"a}, {\'e}, {\^o}, {\v{s}}, ...) -> Unicode, via combining marks and NFC.
        marks = {'"': "\u0308", "'": "\u0301", "`": "\u0300", "^": "\u0302", "~": "\u0303", "v": "\u030c",
                 "c": "\u0327", "=": "\u0304", ".": "\u0307", "u": "\u0306"}
        val = re.sub(r"\\([\"'`^~v=.uc])\s*\{?\s*([A-Za-z])\s*\}?",
                     lambda m: unicodedata.normalize("NFC", m.group(2) + marks[m.group(1)]), val)
        val = val.replace("\\&", "&").replace("--", "-")
        fields[f.group(1).lower()] = re.sub(r"[{}]", "", val)
    entries[m.group(2).strip()] = fields
order: list[str] = []


def cite(match):
    nums = []
    for k in (k.strip().lstrip("@") for k in match.group(1).split(";")):
        if k not in order:
            order.append(k)
        nums.append(order.index(k) + 1)
    return "[" + ", ".join(str(n) for n in sorted(nums)) + "]"


def fmt_ref(k):
    e = entries.get(k, {})
    names = []
    for a in e.get("author", "").split(" and "):
        a = a.strip()
        if not a:
            continue
        last, first = (a.split(",", 1) + [""])[:2] if "," in a else (a.split()[-1], " ".join(a.split()[:-1]))
        initials = "".join(w[0] for w in re.split(r"[\s-]+", first.strip()) if w)
        names.append(f"{last.strip()} {initials}".strip())
    who = ", ".join(names[:6]) + (", et al." if len(names) > 6 else "")
    venue = e.get("journal") or e.get("booktitle") or ""
    if not venue and e.get("eprint"):
        venue = "arXiv:" + e["eprint"]
    vol = e.get("volume", "")
    pages = e.get("pages", "").replace("--", "-")
    tail = ", ".join(x for x in (venue, e.get("year", ""), (vol + (f":{pages}" if pages else "")) if vol else pages) if x)
    who = who.rstrip(".")
    title = e.get("title", k).rstrip(".")
    return f"{who}. {title}. {tail}."


def set_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_runs(p, text, size=None):
    text = re.sub(r"\[(@[^\]]+)\]", cite, text)
    text = text.replace("`", "")
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(p.add_run(part[2:-2]), size, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_font(p.add_run(part[1:-1]), size, italic=True)
        else:
            set_font(p.add_run(part), size)


def fmt_par(p, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, before=0, indent=None):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = 1.15
    if indent is not None:
        pf.left_indent = pf.right_indent = Inches(indent)


def border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for edge, sz in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if sz else "nil")
        el.set(qn("w:sz"), str(sz or 0))
        el.set(qn("w:color"), "000000")
        b.append(el)


def page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_font(run, 9)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(1.0)
sec.top_margin = sec.bottom_margin = Inches(1.0)
page_number(sec)
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(11)

lines = md.splitlines()
i, para, section_name = 0, [], ""
in_appendix = False


def flush():
    global para
    if not para:
        return
    text = " ".join(para)
    para = []
    if re.match(r"^\*Figure \d+\.", text):                       # figure caption, below the figure
        body = text.strip("*")
        label, rest = body.split(".", 1)
        p = doc.add_paragraph()
        set_font(p.add_run(label + "."), 9.5, bold=True)
        add_runs(p, rest, 9.5)
        fmt_par(p, 9.5, after=10)
        return
    if re.match(r"^Table [A-Z]?\d+\.", text):                          # table caption, above the table
        label, rest = text.split(".", 1)
        p = doc.add_paragraph()
        set_font(p.add_run(label + "."), 9.5, bold=True)
        add_runs(p, rest, 9.5)
        fmt_par(p, 9.5, after=3, before=6)
        p.paragraph_format.keep_with_next = True
        return
    p = doc.add_paragraph()
    if section_name == "Abstract":
        add_runs(p, text, 10.5)
        fmt_par(p, 10.5, indent=0.4, after=10)
    elif re.match(r"^[A-Z]\w* = .*\(\d+\)$", text):                # displayed equation, numbered
        add_runs(p, text)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=8)
    elif re.match(r"^\d+\. ", text):                              # numbered list item, hanging indent
        add_runs(p, text)
        fmt_par(p, after=3)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.25)
    else:
        add_runs(p, text)
        fmt_par(p)


while i < len(lines):
    ln = lines[i]
    if ln.startswith("# "):
        flush()
        p = doc.add_paragraph()
        set_font(p.add_run(ln[2:].strip()), 16, bold=True)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
        p = doc.add_paragraph()
        set_font(p.add_run(AUTHOR), 11)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
        if AFFIL:
            p = doc.add_paragraph()
            set_font(p.add_run(AFFIL), 10)
            fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
        p = doc.add_paragraph()
        set_font(p.add_run(STATUS), 10, italic=True)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    elif ln.startswith("## ") or ln.startswith("### "):
        flush()
        level = 2 if ln.startswith("## ") else 3
        title = ln.lstrip("#").strip()
        section_name = title
        if title == "References":
            i += 1
            while i < len(lines) and not lines[i].startswith("## "):
                i += 1
            p = doc.add_paragraph()
            set_font(p.add_run("References"), 12, bold=True)
            fmt_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6)
            refs_anchor = p
            continue
        if title.startswith("Appendix") and not in_appendix:
            in_appendix = True
        p = doc.add_paragraph()
        if title == "Abstract":
            set_font(p.add_run("Abstract"), 11, bold=True)
            fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
        else:
            set_font(p.add_run(title), 12 if level == 2 else 11, bold=True, italic=(level == 3))
            fmt_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=12 if level == 2 else 8, after=4)
        p.paragraph_format.keep_with_next = True
    elif ln.startswith("|"):
        flush()
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|[-| ]+\|$", lines[i]):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        ncol = len(rows[0])
        size = 7 if ncol >= 12 else (7.5 if ncol >= 9 else (8.5 if ncol >= 6 else 9.5))
        t = doc.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True
        for r, row in enumerate(rows):
            for c in range(ncol):
                cell = t.cell(r, c)
                cell.text = ""
                p = cell.paragraphs[0]
                add_runs(p, row[c] if c < len(row) else "", size)
                for run in p.runs:
                    run.bold = run.bold or r == 0
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c < (2 if ncol >= 9 else 1) else WD_ALIGN_PARAGRAPH.CENTER
                edges = {}
                if r == 0:
                    edges.update(top=12, bottom=6)
                if r == len(rows) - 1:
                    edges.update(bottom=12)
                edges.setdefault("left", 0)
                edges.setdefault("right", 0)
                border(cell, **edges)
        # Explicit column widths (text columns wider) and a table that never splits across pages.
        # Width follows the longest body entry (header words may wrap), so number columns such as the CI never wrap.
        total, cw, pad = 6.5, size * 0.55 / 72, (0.12 if ncol >= 12 else 0.17)  # inches per character at this font size; cell padding
        need = [max(max((len(r[c]) for r in rows[1:] if c < len(r)), default=4),
                    max((len(w) for w in rows[0][c].split()), default=4)) * cw + pad for c in range(ncol)]
        for c in range(ncol):  # an interval column wraps after 'to', which frees width for the text columns
            body = [r[c] for r in rows[1:] if c < len(r) and r[c]]
            if body and all(re.match(r"^[-+]?[\d.]+ to [-+]?[\d.]+$", b) for b in body):
                need[c] = max(len(b.split(' to ')[0]) + 3 for b in body) * cw + pad
        widths = list(need)
        while sum(widths) > total:  # shrink the widest (text) column until the table fits the text block
            j = max(range(ncol), key=lambda k: widths[k])
            widths[j] -= min(0.05, sum(widths) - total)
        spare = total - sum(widths)
        widths = [w + spare * w / sum(widths) for w in widths]
        t.autofit = False
        if ncol >= 12:  # narrow cell margins (0.04 in) so the 0.12 in padding above is what Word uses
            mar = OxmlElement("w:tblCellMar")
            for side in ("left", "right"):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:w"), "58")
                e.set(qn("w:type"), "dxa")
                mar.append(e)
            t._tbl.tblPr.append(mar)
        for c in range(ncol):  # Word lays out from the grid columns, not only the cell widths
            t.columns[c].width = Inches(widths[c])
        for r, row_obj in enumerate(t.rows):
            trPr = row_obj._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            for c, cell in enumerate(row_obj.cells):
                cell.width = Inches(widths[c])
                if r < len(t.rows) - 1:
                    cell.paragraphs[0].paragraph_format.keep_with_next = True
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        continue
    elif ln.startswith("!["):
        flush()
        path = re.search(r"\(([^)]+)\)", ln).group(1)
        doc.add_picture(str(SRC / path), width=Inches(6.4 if ("fig0" in path or "fig1" in path) else 3.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    elif not ln.strip():
        flush()
    elif re.match(r"^\d+\. ", ln):                                   # a new numbered item starts its own paragraph
        flush()
        para.append(ln.strip())
    else:
        para.append(ln.strip())
    i += 1
flush()

# References are placed where the "References" heading was, i.e. before the appendices.
anchor = refs_anchor
for n, k in reversed(list(enumerate(order, 1))):
    p = doc.add_paragraph()
    set_font(p.add_run(f"[{n}] "), 9.5)
    set_font(p.add_run(fmt_ref(k)), 9.5)
    fmt_par(p, 9.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    anchor._p.addnext(p._p)
out = Path([a for a in sys.argv[1:] if not a.startswith("--")][0])
doc.save(out)
missing = [k for k in order if k not in entries]
print("saved", out, "refs", len(order), "missing", missing)
