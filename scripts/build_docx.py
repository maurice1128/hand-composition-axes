#!/usr/bin/env python
"""Render docs/PAPER_RAL.md to a Word file laid out like the paper.

For circulating the manuscript to readers who will not open LaTeX. The
submission artifact is `docs/submission/`; this is a reading copy.

The layout follows the compiled paper: a one-column title block over a
two-column body, IEEE-ish headings, tables sized to the column, and footnotes
collected at the end with superscript markers where they occur.

**The provenance line is stripped by matching that line only.** An earlier
version matched from it to the next horizontal rule with DOTALL, and since the
draft puts a rule after the abstract, the whole abstract silently vanished from
every Word copy produced. The line-anchored pattern below cannot do that, and
the abstract is asserted present before the file is written.

Usage:
    python scripts/build_docx.py [--md docs/PAPER_RAL.md] [--out docs/paper.docx]
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent

#: Inline markup, resolved in one pass. Code first, so markup characters inside
#: a code span are not read as markup.
INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[\^[\w-]+\])")


def set_columns(section, n: int) -> None:
    """Two-column body. python-docx exposes no API for this."""
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "360")


def add_runs(par, text: str, notes: dict[str, int]) -> None:
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("[^") and piece.endswith("]"):
            key = piece[2:-1]
            run = par.add_run(str(notes.get(key, "?")))
            run.font.superscript = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = par.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8)
        elif piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def add_table(doc: Document, rows: list[str], notes: dict[str, int]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], cells[2:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = True
    for i, h in enumerate(header):
        add_runs(t.rows[0].cells[i].paragraphs[0], h, notes)
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in body:
        row = (row + [""] * len(header))[:len(header)]
        for i, v in enumerate(row):
            add_runs(t.add_row().cells[i].paragraphs[0] if i == 0
                     else t.rows[-1].cells[i].paragraphs[0], v, notes)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(7)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--md", default="docs/PAPER_RAL.md")
    ap.add_argument("--out", default="docs/paper.docx")
    args = ap.parse_args()

    md = (ROOT / args.md).read_text(encoding="utf-8").replace("\r\n", "\n")
    title = md.splitlines()[0].lstrip("# ").strip()
    # That one line, and only that line.
    md = re.sub(r"^\*\*Target:[^\n]*\n", "", md, flags=re.M)
    md = re.sub(r"^# .*\n", "", md, count=1, flags=re.M)

    # Number the footnotes in order of definition so markers can resolve.
    notes: dict[str, int] = {}
    bodies: dict[str, str] = {}
    for m in re.finditer(r"^\[\^([\w-]+)\]:[ \t]*((?:[^\n]*\n?)(?:(?![ \t]*$|\[\^)[^\n]*\n?)*)",
                         md, flags=re.M):
        notes[m.group(1)] = len(notes) + 1
        bodies[m.group(1)] = " ".join(m.group(2).split())
    md = re.sub(r"^\[\^[\w-]+\]:.*?(?=\n[ \t]*\n)", "", md, flags=re.M | re.S)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9.5)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.75)
        s.left_margin = s.right_margin = Inches(0.7)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(title)
    run.bold = True
    run.font.size = Pt(17)
    a = doc.add_paragraph()
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a.add_run("Mu-Hua Wang\nNational Yang Ming Chiao Tung University\n"
              "mauricewang1128@gmail.com").font.size = Pt(9.5)

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    body.top_margin = body.bottom_margin = Inches(0.75)
    body.left_margin = body.right_margin = Inches(0.7)
    set_columns(body, 2)

    lines = md.split("\n")
    i = 0
    para: list[str] = []
    n_tables = n_head = n_para = n_figs = 0

    def flush() -> None:
        nonlocal para, n_para
        if para:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, " ".join(" ".join(para).split()), notes)
            n_para += 1
            para = []

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if s.startswith("|") and i + 1 < len(lines) and set(
                lines[i + 1].replace("|", "").replace(":", "").strip()) <= {"-", " "}:
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block, notes)
            n_tables += 1
            continue

        # A figure block, possibly hard-wrapped, ends at its closing brace.
        if s.startswith("!["):
            flush()
            block = [s]
            while i + 1 < len(lines) and not block[-1].rstrip().endswith("}"):
                i += 1
                block.append(lines[i].strip())
            joined = " ".join(" ".join(block).split())
            mf = re.match(r"^!\[(.*)\]\(figures/([^)]+)\)\{#fig:[\w-]+\}$", joined)
            if not mf:
                raise SystemExit("malformed figure block: " + joined[:70])
            cap, name = mf.groups()
            img = ROOT / "docs" / "figures" / name
            if not img.exists():
                raise SystemExit("missing figure: " + str(img))
            doc.add_picture(str(img), width=Inches(3.3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            fc = doc.add_paragraph()
            fc.paragraph_format.space_after = Pt(6)
            add_runs(fc, cap, notes)
            for r in fc.runs:
                r.font.size = Pt(8)
            n_figs += 1
            i += 1
            continue

        m = re.match(r"^(#{2,4})\s+(.*)$", s)
        if m:
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(m.group(2).strip())
            r.bold = True
            r.font.size = Pt(11 if len(m.group(1)) == 2 else 10)
            n_head += 1
            i += 1
            continue

        if re.match(r"^\d+\.\s", s):
            flush()
            item = [s]
            i += 1
            while i < len(lines) and lines[i].startswith("   "):
                item.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, re.sub(r"^\d+\.\s*", "", " ".join(item)), notes)
            continue

        if ln.startswith("    ") and s:
            flush()
            p = doc.add_paragraph()
            r = p.add_run(ln.strip())
            r.font.name = "Consolas"
            r.font.size = Pt(8)
            i += 1
            continue

        if s in ("---", "***") or not s:
            flush()
            i += 1
            continue

        para.append(s)
        i += 1
    flush()

    if bodies:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run("Footnotes")
        r.bold = True
        r.font.size = Pt(10)
        for key, num in sorted(notes.items(), key=lambda kv: kv[1]):
            fp = doc.add_paragraph()
            fp.paragraph_format.space_after = Pt(2)
            fp.add_run(f"{num}  ").font.superscript = True
            add_runs(fp, bodies[key], notes)
            for run in fp.runs:
                run.font.size = Pt(8)

    text = "\n".join(p.text for p in doc.paragraphs)
    if "Modular motion priors decompose" not in text:
        raise SystemExit("the abstract is missing from the output; refusing to write")
    if "[^" in text:
        raise SystemExit("an unresolved footnote marker reached the output")

    out = ROOT / args.out
    doc.save(out)
    print(f"wrote {out}")
    print(f"  {n_head} headings, {n_para} paragraphs, {n_tables} tables, "
          f"{n_figs} figures, "
          f"{len(bodies)} footnotes, two-column body")
    print("  Reading copy. docs/submission/ is the submission artifact.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
